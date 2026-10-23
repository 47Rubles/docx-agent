from docx.oxml import OxmlElement
from docx.oxml.ns import qn

COLORS = {'info':('2563EB','EFF6FF'),'warning':('D97706','FFFBEB'),'success':('059669','ECFDF5'),'critical':('D32F2F','FDF2F2')}

def build_callout(doc, block):
    accent, bg = COLORS[block.callout_type]
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0,0)
    cell.text = ''
    title = cell.paragraphs[0].add_run(block.title)
    title.bold = True
    cell.add_paragraph(block.text)
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), bg); tcpr.append(shd)
    borders = OxmlElement('w:tcBorders')
    for side in ('left','top','right','bottom'):
        el = OxmlElement('w:'+side)
        el.set(qn('w:val'),'single' if side=='left' else 'nil')
        if side=='left':
            el.set(qn('w:sz'),'24') # 3 pt; w:sz is measured in eighths of a point
            el.set(qn('w:color'),accent)
        borders.append(el)
    tcpr.append(borders)
    margins = OxmlElement('w:tcMar')
    for side in ('top','left','bottom','right'):
        el = OxmlElement('w:'+side)
        el.set(qn('w:w'),'160'); el.set(qn('w:type'),'dxa')
        margins.append(el)
    tcpr.append(margins)
    doc.add_paragraph()
    return table
