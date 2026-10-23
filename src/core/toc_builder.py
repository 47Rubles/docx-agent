"""Shared native Word field helpers."""
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def add_field(paragraph, instruction, placeholder='1'):
    for typ in ('begin',):
        node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),typ);paragraph.add_run()._r.append(node)
    instr=OxmlElement('w:instrText');instr.set('{http://www.w3.org/XML/1998/namespace}space','preserve');instr.text=' '+instruction+' '
    paragraph.add_run()._r.append(instr)
    for typ in ('separate',):
        node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),typ);paragraph.add_run()._r.append(node)
    paragraph.add_run(placeholder)
    node=OxmlElement('w:fldChar');node.set(qn('w:fldCharType'),'end');paragraph.add_run()._r.append(node)
