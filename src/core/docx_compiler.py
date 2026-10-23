from pathlib import Path
import logging
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from src.models.ast_schema import DocumentAST, ParagraphBlock, FormulaBlock, TableBlock, CalloutBlock
from src.core.omml_converter import latex_to_omml
from src.core.table_builder import build_table
from src.core.callout_builder import build_callout

LOG = logging.getLogger(__name__)
TEMPLATE = Path(__file__).resolve().parents[2]/'assets/templates/default.docx'

def add_field(paragraph,instruction,placeholder='1'):
    for kind in ('begin',):
        r=paragraph.add_run()._r; el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),kind); r.append(el)
    r=paragraph.add_run()._r; el=OxmlElement('w:instrText'); el.set('{http://www.w3.org/XML/1998/namespace}space','preserve'); el.text=f' {instruction} '; r.append(el)
    r=paragraph.add_run()._r; el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),'separate'); r.append(el)
    paragraph.add_run(placeholder)
    r=paragraph.add_run()._r; el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),'end'); r.append(el)

def compile_document(ast:DocumentAST,output:str|Path,template:str|Path|None=None)->Path:
    path=Path(template) if template else TEMPLATE
    doc=Document(str(path)) if path.exists() else Document()
    meta=ast.metadata
    doc.core_properties.title=meta.title; doc.core_properties.author=meta.author
    doc.add_heading(meta.title,0)
    if meta.subtitle: doc.add_paragraph(meta.subtitle,style='Subtitle')
    if meta.author: doc.add_paragraph(meta.author,style='Subtitle')
    if meta.enable_toc:
        doc.add_paragraph('Оглавление',style='Title')
        add_field(doc.add_paragraph(),r'TOC \o "1-3" \h \z \u','Обновите оглавление в Word: Ctrl+A → F9')
        doc.add_page_break()
    for section in ast.sections:
        doc.add_heading(section.title,level=section.level)
        for block in section.blocks:
            try:
                if isinstance(block,ParagraphBlock):
                    p=doc.add_paragraph(style=block.style if block.style in doc.styles else 'Normal')
                    if block.runs:
                        for part in block.runs:
                            run=p.add_run(part.text); run.bold=part.bold; run.italic=part.italic
                            if part.code: run.font.name='Consolas'
                    else:
                        run=p.add_run(block.text); run.bold=block.bold; run.italic=block.italic
                elif isinstance(block,FormulaBlock):
                    p=doc.add_paragraph()
                    if block.display_mode=='block': p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    try: p._p.append(latex_to_omml(block.latex,block.display_mode=='block'))
                    except Exception as exc:
                        LOG.warning('Formula conversion failed: %s',exc)
                        run=p.add_run('[Формула не преобразована] '+block.latex);run.italic=True;run.font.name='Consolas'
                elif isinstance(block,TableBlock): build_table(doc,block)
                elif isinstance(block,CalloutBlock): build_callout(doc,block)
            except Exception as exc:
                LOG.exception('Failed to compile block in %s',section.id)
                doc.add_paragraph('[Ошибка блока: '+type(block).__name__+'] '+str(exc))
    footer=doc.sections[0].footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Страница ');add_field(footer,'PAGE');footer.add_run(' из ');add_field(footer,'NUMPAGES')
    settings=doc.settings.element
    update=settings.find(qn('w:updateFields'))
    if update is None: update=OxmlElement('w:updateFields');settings.append(update)
    update.set(qn('w:val'),'true')
    if meta.enable_track_changes:
        tracking=settings.find(qn('w:trackRevisions'))
        if tracking is None: tracking=OxmlElement('w:trackRevisions'); settings.append(tracking)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(str(output))
    return output
